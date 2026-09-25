"""Excel / CSV export and Google Sheets sync."""
from __future__ import annotations

import csv
import json
from pathlib import Path
from typing import Dict, Iterable

from .settings import env, resolve_path
from .textutil import clean_cell

COLUMNS = [
    ("score", "Score", 8), ("title", "Title", 45), ("company", "Company", 24), ("location", "Location", 22),
    ("category", "Category", 16), ("job_type", "Type", 14), ("salary", "Salary", 20), ("source", "Source", 22),
    ("status", "Status", 12), ("found_at", "Found", 20), ("posted_at", "Posted", 20), ("reasons", "Why", 50),
    ("notes", "Notes", 30), ("url", "Link", 40),
]


def _value(row: dict, key: str):
    value = row.get(key)
    if key == "reasons":
        try:
            return " | ".join(json.loads(value or "[]"))
        except ValueError:
            return value or ""
    if key in ("found_at", "posted_at") and value:
        return str(value)[:16].replace("T", " ")
    return value if value is not None else ""


def export_xlsx(path: Path, sheets: Dict[str, Iterable]) -> Path:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    header_fill = PatternFill("solid", fgColor="123A5E")
    fills = {
        "hi": PatternFill("solid", fgColor="D8F3DC"),
        "mid": PatternFill("solid", fgColor="FFF3C4"),
    }
    wb = Workbook()
    wb.remove(wb.active)
    for name, rows in sheets.items():
        ws = wb.create_sheet(title=name[:31])
        ws.append([h for _, h, _ in COLUMNS])
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = header_fill
            cell.alignment = Alignment(vertical="center")
        for raw in rows:
            row = dict(raw)
            ws.append([clean_cell(_value(row, key)) if key != "score" else row.get("score") or 0 for key, _, _ in COLUMNS])
            r = ws.max_row
            link = ws.cell(row=r, column=len(COLUMNS))
            if row.get("url"):
                link.hyperlink = row["url"]
                link.style = "Hyperlink"
            score = row.get("score") or 0
            fill = fills["hi"] if score >= 70 else fills["mid"] if score >= 50 else None
            if fill:
                ws.cell(row=r, column=1).fill = fill
        for idx, (_, _, width) in enumerate(COLUMNS, start=1):
            ws.column_dimensions[ws.cell(row=1, column=idx).column_letter].width = width
        ws.freeze_panes = "C2"
        ws.auto_filter.ref = ws.dimensions
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


def export_csv(path: Path, rows: Iterable) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    with open(path, "w", newline="", encoding="utf-8-sig") as fh:
        writer = csv.writer(fh)
        writer.writerow([h for _, h, _ in COLUMNS])
        for raw in rows:
            row = dict(raw)
            writer.writerow([_value(row, key) for key, _, _ in COLUMNS])
    return path


# ---------------------------------------------------------------- Google Sheets
SHEET_HEADER = ["ID", "Found", "Score", "Title", "Company", "Location", "Category", "Type", "Salary",
                "Source", "Status", "Notes", "Why", "Link"]


def sheet_configured() -> bool:
    return bool(env("GOOGLE_SHEET_ID") and (env("GOOGLE_SERVICE_ACCOUNT_JSON") or env("GOOGLE_SERVICE_ACCOUNT_FILE")))


def _safe(value) -> str:
    text = clean_cell(value)
    return "'" + text if text[:1] in ("=", "+", "-", "@") else text


def sync_google_sheet(con, min_score: int) -> str:
    """Append new opportunities (score >= min_score) to the sheet. Existing rows —
    and the Status/Notes you edit in the sheet — are never overwritten."""
    if not sheet_configured():
        return "skipped (not configured)"
    import gspread

    info = env("GOOGLE_SERVICE_ACCOUNT_JSON")
    if info:
        client = gspread.service_account_from_dict(json.loads(info))
    else:
        client = gspread.service_account(filename=str(resolve_path(env("GOOGLE_SERVICE_ACCOUNT_FILE"))))
    book = client.open_by_key(env("GOOGLE_SHEET_ID"))
    tab = env("GOOGLE_SHEET_TAB", "Opportunities")
    try:
        ws = book.worksheet(tab)
    except gspread.WorksheetNotFound:
        ws = book.add_worksheet(title=tab, rows=1000, cols=len(SHEET_HEADER))

    existing = ws.col_values(1)
    if not existing:
        ws.append_row(SHEET_HEADER, value_input_option="RAW")
        existing = [SHEET_HEADER[0]]
    known = set(existing)

    threshold = int(env("SHEET_MIN_SCORE") or min_score)
    rows = con.execute(
        "SELECT * FROM opportunities WHERE synced=0 AND score>=? ORDER BY score DESC", (threshold,)
    ).fetchall()
    values = []
    for row in rows:
        short_id = row["uid"][:16]
        if short_id in known:
            continue
        url = (row["url"] or "").replace('"', '""')
        values.append([
            short_id, _value(dict(row), "found_at"), row["score"], _safe(row["title"]), _safe(row["company"]),
            _safe(row["location"]), row["category"], _safe(row["job_type"]), _safe(row["salary"]),
            _safe(row["source"]), row["status"], _safe(row["notes"]), _safe(_value(dict(row), "reasons")),
            f'=HYPERLINK("{url}","Open")' if url else "",
        ])
    if values:
        ws.append_rows(values, value_input_option="USER_ENTERED")
    if rows:
        con.executemany("UPDATE opportunities SET synced=1 WHERE id=?", [(r["id"],) for r in rows])
        con.commit()
    return f"sent ({len(values)} rows)"
